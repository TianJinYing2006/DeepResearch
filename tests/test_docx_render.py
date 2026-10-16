"""需求 38 F3：DOCX 渲染器单测（结构等价断言；零网络 / 零 PG）。

断言**结构**（样式名/层级/文本）而非字节——python-docx/模板可演进，
不承诺跨版本二进制一致（需求 38 §9）。
"""
from __future__ import annotations

import io

from docx import Document

from web.backend.docx_render import render_docx
from web.backend.export import ExportDocument, document_from_payload


def _payload(**overrides) -> dict:
    payload = {
        "run_id": "run-docx", "topic": "DOCX 渲染主题",
        "generated_at": "2026-10-16T00:00:00+00:00", "run_status": "success",
        "stop_reason": "completed", "cancelled": False, "token_used": 100,
        "cost_estimate_cny": 0.001, "elapsed_seconds": 12.5,
        "degradation_count": 1, "depth": 2,
        "result": {
            "report": "# 报告标题\n\n中文正文段落。\n\n## 小节\n\n- 要点一\n- 要点二"
                      "\n\n1. 甲\n2. 乙\n\n> 引用\n\n```\ncode line\n```",
            "citations": [{"claim": "论断", "source": "https://e.com/a",
                           "source_type": "web", "confidence": 0.9,
                           "existence": True, "supported": True}],
            "visited_sources": ["https://e.com/a"],
        },
    }
    payload.update(overrides)
    return payload


def _render(payload: dict):
    data = render_docx(document_from_payload(payload))
    assert data[:2] == b"PK"  # docx = zip 容器
    return data, Document(io.BytesIO(data))


def _styled(document) -> list:
    return [(p.style.name, p.text) for p in document.paragraphs]


def test_docx_structure_matches_markdown_semantics():
    _, document = _render(_payload())
    styled = _styled(document)
    texts = [t for _, t in styled]

    # 头部：标题 + 三行元数据（与 render_markdown 同信息）
    assert styled[0][0] == "Title" and "DOCX 渲染主题：DeepResearch 报告导出" in styled[0][1]
    assert any("run_status success" in t and "停止原因 completed" in t for t in texts)
    assert any("token 100" in t and "降级条目 1" in t and "检索深度 2" in t for t in texts)

    # 正文结构：标题层级 / 段落 / 列表 / 引用 / 代码块
    assert ("Heading 1", "报告标题") in styled
    assert ("Heading 2", "小节") in styled
    assert ("Normal", "中文正文段落。") in styled
    assert ("List Bullet", "要点一") in styled and ("List Bullet", "要点二") in styled
    assert ("List Number", "甲") in styled and ("List Number", "乙") in styled
    assert ("Quote", "引用") in styled
    assert any(style == "Normal" and t == "code line" for style, t in styled)

    # 引用清单 / 访问来源（与 markdown 同节名与条目文本）
    assert ("Heading 2", "引用清单") in styled
    assert any(style == "List Number" and "论断 — <https://e.com/a>" in t
               and "置信度 0.90" in t and "存在性校验通过" in t for style, t in styled)
    assert ("Heading 2", "访问过的来源") in styled
    assert ("List Bullet", "<https://e.com/a>") in styled


def test_docx_unavailable_and_empty_semantics():
    # citations / visited_sources = None ⇒ 显式「不可用」（不伪造空清单）。
    # 注意：`document_from_payload` 会把 None 归一为 []（内存路径字段完整，无不可用态）
    # ⇒ 这里直接构造 ExportDocument（等价 store 降级路径的产物）。
    doc = ExportDocument(run_id="run-docx", topic="DOCX 渲染主题",
                         generated_at="2026-10-16T00:00:00+00:00", run_status="success",
                         report="# 正文", citations=None, visited_sources=None)
    document = Document(io.BytesIO(render_docx(doc)))
    texts = [t for _, t in _styled(document)]
    assert "（不可用：历史产物未保留引用数据）" in texts
    assert "（不可用：历史产物未保留来源数据）" in texts

    # 空列表 ⇒ 省略整节（与既有 markdown 口径一致）
    payload2 = _payload()
    payload2["result"]["citations"] = []
    payload2["result"]["visited_sources"] = []
    _, document2 = _render(payload2)
    texts2 = [t for _, t in _styled(document2)]
    assert "引用清单" not in texts2 and "访问过的来源" not in texts2


def test_docx_empty_report_explicit():
    payload = _payload()
    payload["result"]["report"] = ""
    _, document = _render(payload)
    assert "（本次运行未产出报告正文）" in [t for _, t in _styled(document)]


def test_docx_cjk_font_pinned():
    """东亚字体钉到 Normal 样式（Word 取字体路径 w:eastAsia）。"""
    data, _ = _render(_payload())
    import zipfile

    from docx.oxml.ns import qn

    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        styles_xml = zf.read("word/styles.xml").decode("utf-8")
    assert "w:eastAsia" in styles_xml and "宋体" in styles_xml
    # qn 引用仅用于自证命名空间前缀存在（防常量被误删）
    assert qn("w:eastAsia") == "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia"
