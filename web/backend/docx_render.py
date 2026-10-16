"""需求 38 F3：DOCX 渲染器（python-docx **写作侧**；同步、轻量、离线确定性）。

与 `export.render_markdown` **同一信息结构**（唯一数据源 = `ExportDocument`）：
元数据头 + 报告正文（标题层级/段落/列表/引用/代码块）+ 引用清单 + 访问过的来源；
历史缺失字段显式「不可用」（不伪造空清单）。

边界（需求 38 §12.2 与批次纪律）：
- 报告正文**原文照录**：只按行结构转排（标题/列表/引用/代码），文本逐字保留；
- 不落盘、不写库（按需生成、内存返回，D-19 延续）；
- 确定性：无网络、无外部服务；不承诺跨版本字节一致（python-docx/模板可演进），
  测试断言**结构**（样式名/文本/层级）而非二进制。
"""
from __future__ import annotations

import io
import re
from typing import Any, List, Optional

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .export import ExportDocument

#: 东亚字体（Word 通用字面名；阅读器缺字体时自行替换，不影响文本层）
_CJK_FONT = "宋体"
#: 代码块西文字体（等宽）
_MONO_FONT = "Consolas"

_ORDERED_RE = re.compile(r"^\d{1,4}\.\s+(.*)$")


def _set_cjk(style: Any) -> None:
    """把样式的东亚字体钉到 `_CJK_FONT`（rFonts w:eastAsia；Word 对 CJK 的取字体路径）。"""
    style.font.name = style.font.name or "Calibri"
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), _CJK_FONT)


def _add_meta(document: Any, text: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def _add_body(document: Any, report: str) -> None:
    """报告正文：受限行结构转排（标题/无序/有序/引用/代码块/段落），文本逐字保留。"""
    para: List[str] = []
    in_code = False
    code_lines: List[str] = []

    def flush_para() -> None:
        if para:
            document.add_paragraph("\n".join(para))
            para.clear()

    def flush_code() -> None:
        if code_lines:
            paragraph = document.add_paragraph()
            run = paragraph.add_run("\n".join(code_lines))
            run.font.name = _MONO_FONT
            run.font.size = Pt(9.5)
            code_lines.clear()

    for raw in (report or "").splitlines():
        line = raw.rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            flush_para()
            if in_code:
                flush_code()
            in_code = not in_code
            continue
        if in_code:
            code_lines.append(line)
            continue
        if not stripped:
            flush_para()
            continue
        if stripped.startswith("#"):
            flush_para()
            level = min(6, len(stripped) - len(stripped.lstrip("#")))
            document.add_heading(stripped[level:].strip(), level=level)
            continue
        if stripped.startswith(("- ", "* ")):
            flush_para()
            document.add_paragraph(stripped[2:].strip(), style="List Bullet")
            continue
        match = _ORDERED_RE.match(stripped)
        if match:
            flush_para()
            document.add_paragraph(match.group(1).strip(), style="List Number")
            continue
        if stripped.startswith(">"):
            flush_para()
            document.add_paragraph(stripped.lstrip(">").strip(), style="Quote")
            continue
        para.append(line)
    flush_para()
    flush_code()


def _citation_text(citation: dict) -> str:
    verified = "存在性校验通过" if citation.get("existence") else "存在性校验未通过"
    supported = "论断受支持" if citation.get("supported") else "论断未获支持"
    confidence = citation.get("confidence")
    conf = f"{confidence:.2f}" if isinstance(confidence, (int, float)) else "n/a"
    claim = str(citation.get("claim", "")).strip() or "（无论断）"
    note = str(citation.get("note", "")).strip()
    note = f"　备注：{note}" if note else ""
    return (f"{claim} — <{citation.get('source', '')}>"
            f"（{citation.get('source_type', 'unknown')} · 置信度 {conf}"
            f" · {verified} · {supported}）{note}")


def _meta_value(value: Any) -> Any:
    """不可知字段（None）显式渲染为 n/a（不伪造 0）——与 `render_markdown` 同口径。"""
    return "n/a" if value is None else value


def render_docx(doc: ExportDocument) -> bytes:
    """`ExportDocument` → DOCX 字节（结构等价于 `render_markdown` 的信息结构）。"""
    document = Document()
    _set_cjk(document.styles["Normal"])
    document.styles["Normal"].font.size = Pt(11)

    document.add_heading(f"{doc.topic}：DeepResearch 报告导出", level=0)
    _add_meta(document, f"run_id {doc.run_id} · 导出时间 {doc.generated_at}")
    _add_meta(document, f"run_status {doc.run_status} · 停止原因 {doc.stop_reason}"
                       f"{' · 已取消' if doc.cancelled else ''}")
    _add_meta(document, f"token {doc.token_used} · 成本估算 ¥{doc.cost_estimate_cny}"
                       f" · 耗时 {doc.elapsed_seconds}s"
                       f" · 降级条目 {_meta_value(doc.degradation_count)}"
                       f" · 检索深度 {_meta_value(doc.depth)}")

    report = doc.report.strip()
    if report:
        _add_body(document, report)
    else:
        document.add_paragraph("（本次运行未产出报告正文）")

    citations: Optional[List[dict]] = doc.citations
    if citations is None:
        document.add_heading("引用清单", level=2)
        document.add_paragraph("（不可用：历史产物未保留引用数据）")
    elif citations:
        document.add_heading("引用清单", level=2)
        for citation in citations:
            document.add_paragraph(_citation_text(citation), style="List Number")

    visited = doc.visited_sources
    if visited is None:
        document.add_heading("访问过的来源", level=2)
        document.add_paragraph("（不可用：历史产物未保留来源数据）")
    elif visited:
        document.add_heading("访问过的来源", level=2)
        for source in visited:
            document.add_paragraph(f"<{source}>", style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
